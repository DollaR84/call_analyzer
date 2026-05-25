from datetime import datetime
import logging
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from schemas import Transcript


logger = logging.getLogger(__name__)


class ExcelWriter:

    def __init__(self, data_dir: Path):
        self.data_dir = data_dir

        self.wb: Optional[Workbook] = None
        self.sheet: Optional[Worksheet] = None

    def open(self) -> Optional[Path]:
        xlsx_file = next(self.data_dir.glob("*.xlsx"), None)
        if not xlsx_file:
            logger.warning("No .xlsx files found in '%s'", str(self.data_dir))
            return None

        try:
            self.wb = openpyxl.load_workbook(filename=xlsx_file)
            self.sheet = self.wb[self.wb.sheetnames[0]]
            return xlsx_file
        except OSError as e:
            logger.error("System I/O error while open '%s': %s", xlsx_file.name, e, exc_info=True)
            return None

    def close(self) -> None:
        if self.wb is not None:
            self.wb.close()

    def _find_real_max_row(self, sheet: Worksheet, column: int) -> int:
        cols = list(sheet.iter_cols(min_col=column, max_col=column, min_row=3, max_row=sheet.max_row))
        if not cols or not cols[0]:
            return 2

        cells = cols[0]
        for index in range(len(cells) - 1, -1, -1):
            cell = cells[index]
            if type(cell).__name__ == "MergedCell":
                continue

            if cell.value is not None:
                return index + 3
        return 2

    def _parse_file_datetime(self, file_name: str) -> datetime:
        clean_name = file_name.lstrip(",")
        date_part, time_part, *_ = clean_name.split("_")
        time_clean = time_part.replace("-", ":")
        return datetime.strptime(f"{date_part} {time_clean}", "%Y-%m-%d %H:%M")

    def save(self, transcript: Transcript) -> None:
        xlsx_file = self.open()
        if xlsx_file is None or self.wb is None or self.sheet is None:
            logger.error("Cannot save transcript: Excel file or worksheet is not initialized.")
            return

        column_date = 1
        column_transcript = 21

        real_max_row = self._find_real_max_row(self.sheet, column_date)
        target_dt = self._parse_file_datetime(transcript.file_name)

        current_header = self.sheet.cell(row=2, column=column_transcript).value
        if current_header is None:
            self.sheet.cell(row=2, column=column_transcript, value="transcript")

        target_row = None
        for row in range(3, real_max_row + 1):
            cell_value = self.sheet.cell(row=row, column=column_date).value

            if isinstance(cell_value, datetime):
                if cell_value.replace(second=0, microsecond=0) == target_dt.replace(second=0, microsecond=0):
                    target_row = row
                    break

        if target_row:
            cell = self.sheet.cell(row=target_row, column=column_transcript)
            if type(cell).__name__ != "MergedCell":
                cell.value = transcript.full_text

        else:
            new_row = real_max_row + 1
            self.sheet.cell(row=new_row, column=column_date, value=target_dt)
            cell = self.sheet.cell(row=new_row, column=column_transcript)

            if type(cell).__name__ != "MergedCell":
                cell.value = transcript.full_text

        try:
            self.wb.save(xlsx_file)
            logger.info("saved to excel file successfully: %s", transcript.file_name)
        except PermissionError as e:
            logger.error(
                "permission error: file '%s' is open in another program (Excel) or locked: %s",
                xlsx_file.name, e
            )
        except OSError as e:
            logger.error("System I/O error while saving '%s': %s", xlsx_file.name, e, exc_info=True)

        finally:
            self.close()
            self.sheet = None
            self.wb = None
