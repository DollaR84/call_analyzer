from datetime import datetime
import logging
from pathlib import Path
import threading
from typing import Optional, overload

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from schemas import Report, Transcript

from .base import BaseWriter
from .types import ExcelColumns, ExcelRows


logger = logging.getLogger(__name__)


class ExcelWriter(BaseWriter):

    def __init__(self, output_dir: Path):
        super().__init__(output_dir)

        self._rows_map: dict[str, int] = {}
        self._lock = threading.Lock()

        self.wb: Optional[Workbook] = None
        self.sheet: Optional[Worksheet] = None
        self._red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    def open(self) -> Optional[Path]:
        xlsx_file = next(self.output_dir.glob("*.xlsx"), None)
        if not xlsx_file:
            logger.warning("No .xlsx files found in '%s'", str(self.output_dir))
            return None

        try:
            self.wb = openpyxl.load_workbook(filename=xlsx_file)
            self.sheet = self.wb[self.wb.sheetnames[0]]
            return xlsx_file
        except OSError as e:
            logger.error("System I/O error while open '%s': %s", xlsx_file.name, e, exc_info=True)
            return None

    def close(self) -> None:
        if self.wb is not None:
            self.wb.close()

    def _save_file(self, xlsx_file: Path, transcript_name: str) -> None:
        if self.wb is None:
            return

        try:
            self.wb.save(xlsx_file)
            logger.info("saved to excel file successfully: %s", transcript_name)
        except PermissionError as e:
            logger.error(
                "permission error: file '%s' is open in another program (Excel) or locked: %s",
                xlsx_file.name, e, exc_info=True
            )
        except OSError as e:
            logger.error("System I/O error while saving '%s': %s", xlsx_file.name, e, exc_info=True)

        finally:
            self.close()
            self.sheet = None
            self.wb = None

    def _find_real_max_row(self, sheet: Worksheet, column: int) -> int:
        cols = list(sheet.iter_cols(
            min_col=column,
            max_col=column,
            min_row=ExcelRows.HEADER + 1,
            max_row=sheet.max_row
        ))

        if not cols or not cols[0]:
            return ExcelRows.HEADER

        cells = cols[0]
        for index in range(len(cells) - 1, -1, -1):
            cell = cells[index]
            if type(cell).__name__ == "MergedCell":
                continue

            if cell.value is not None:
                return index + ExcelRows.HEADER + 1
        return ExcelRows.HEADER

    def _parse_file_datetime(self, file_name: str) -> datetime:
        clean_name = file_name.lstrip(",")
        date_part, time_part, *_ = clean_name.split("_")
        time_clean = time_part.replace("-", ":")
        return datetime.strptime(f"{date_part} {time_clean}", "%Y-%m-%d %H:%M")

    @overload
    def save(self, data: Transcript) -> Path:
        ...

    @overload
    def save(self, data: Report) -> Path:
        ...

    def save(self, data: Transcript | Report) -> Path:
        with self._lock:
            if isinstance(data, Transcript):
                return self._save_transcript(data)

            if isinstance(data, Report):
                return self._save_report(data)

        raise TypeError(f"Unsupported data type for excel save: '{type(data).__name__}'")

    def _save_transcript(self, transcript: Transcript) -> Path:
        xlsx_file = self.open()
        if xlsx_file is None or self.wb is None or self.sheet is None:
            raise RuntimeError("Cannot save transcript: Excel file or worksheet is not initialized.")

        real_max_row = self._find_real_max_row(self.sheet, ExcelColumns.DATE)
        target_dt = self._parse_file_datetime(transcript.file_name)
        target_row = None

        for row in range(ExcelRows.HEADER + 1, real_max_row + 1):
            cell_value = self.sheet.cell(row=row, column=ExcelColumns.DATE).value

            if isinstance(cell_value, datetime):
                if cell_value.replace(second=0, microsecond=0) == target_dt.replace(second=0, microsecond=0):
                    target_row = row
                    break

        if target_row:
            cell = self.sheet.cell(row=target_row, column=ExcelColumns.TRANSCRIPT)
            if type(cell).__name__ != "MergedCell":
                cell.value = transcript.full_text

        else:
            target_row = real_max_row + 1
            self.sheet.cell(row=target_row, column=ExcelColumns.DATE, value=target_dt)
            cell = self.sheet.cell(row=target_row, column=ExcelColumns.TRANSCRIPT)

            if type(cell).__name__ != "MergedCell":
                cell.value = transcript.full_text

        self._rows_map[transcript.file_name] = target_row
        self._save_file(xlsx_file, transcript.file_name)
        return xlsx_file

    def _save_report(self, report: Report) -> Path:
        xlsx_file = self.open()
        if xlsx_file is None or self.wb is None or self.sheet is None:
            raise RuntimeError("Cannot save report: Excel file or worksheet is not initialized.")

        target_row = self._rows_map.pop(report.file_name, None)
        if target_row is None:
            raise RuntimeError(
                f"The row for call '{report.file_name}' was not found in Excel "
                "or the report has already been recorded!"
            )

        self.sheet.cell(row=target_row, column=ExcelColumns.REQUESTED_WORK, value=report.requested_work)
        self.sheet.cell(row=target_row, column=ExcelColumns.MANAGER_SCORE, value=report.manager_score)

        cell = self.sheet.cell(row=target_row, column=ExcelColumns.COMMENT)
        cell.value = report.comment
        if report.red_flag:
            cell.fill = self._red_fill

        self.sheet.cell(row=target_row, column=ExcelColumns.PROPER_GREETING, value=report.checklist.proper_greeting)
        self.sheet.cell(row=target_row, column=ExcelColumns.KNOWS_BODY_TYPE, value=report.checklist.knows_body_type)
        self.sheet.cell(row=target_row, column=ExcelColumns.KNOWS_YEAR, value=report.checklist.knows_year)
        self.sheet.cell(row=target_row, column=ExcelColumns.KNOWS_MILEAGE, value=report.checklist.knows_mileage)
        self.sheet.cell(
            row=target_row,
            column=ExcelColumns.OFFERED_COMPLEX_DIAGNOSTICS,
            value=report.checklist.offered_complex_diagnostics
        )

        self.sheet.cell(row=target_row, column=ExcelColumns.KNOWS_HISTORY, value=report.checklist.knows_history)
        self.sheet.cell(row=target_row, column=ExcelColumns.PROPER_GOODBYE, value=report.checklist.proper_goodbye)
        self.sheet.cell(
            row=target_row,
            column=ExcelColumns.FOLLOWED_TOP_100_RULES,
            value=report.checklist.followed_top_100_rules
        )

        self._save_file(xlsx_file, report.file_name)
        return xlsx_file
